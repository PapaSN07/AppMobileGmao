"""
Convert Excel workbook (.xlsx) sheets to JSON files expected by import_to_docker.py.

Usage:
  python scripts/convert_excel_to_json.py --input "C:\\path\\to\\export.xlsx"

Optional:
  python scripts/convert_excel_to_json.py --input "C:\\path\\to\\export.xlsx" --sheet equipment
  python scripts/convert_excel_to_json.py --input "C:\\path\\to\\export.xlsx" --out-dir scripts/data_extracted

Notes:
  - First row in each sheet must contain column names.
  - Empty trailing rows are ignored.
  - Unknown sheet names are exported using <sheet_name>.json (normalized).
"""

from __future__ import annotations

import argparse
import csv
import json
import re
from pathlib import Path
from typing import Dict, Iterable, List, Optional


def normalize_name(value: str) -> str:
    value = value.strip().lower()
    value = re.sub(r"[^a-z0-9]+", "_", value)
    value = re.sub(r"_+", "_", value).strip("_")
    return value


SHEET_TO_OUTPUT: Dict[str, str] = {
    "equipment": "equipment_latest.json",
    "equipement": "equipment_latest.json",

    "attribute": "attributes_latest.json",
    "attributes": "attributes_latest.json",

    "attribute_values": "attribute_values_latest.json",
    "attributevalue": "attribute_values_latest.json",
    "attributevalues": "attribute_values_latest.json",

    "equipment_attribute": "equipment_attribute_latest.json",
    "equipment_attributes": "equipment_attribute_latest.json",

    "equipment_specs": "equipment_specs_latest.json",
    "equipment_specs_": "equipment_specs_latest.json",
    "equipmentspecs": "equipment_specs_latest.json",

    "category_specification": "category_specification_latest.json",
    "category_specifications": "category_specification_latest.json",

    "specification": "specification_latest.json",
    "t_specification": "specification_latest.json",

    "work_order": "ot_workorders_latest.json",
    "work_orders": "ot_workorders_latest.json",
    "workorder": "ot_workorders_latest.json",
    "workorders": "ot_workorders_latest.json",
    "ot": "ot_workorders_latest.json",

    "work_request": "di_workrequests_latest.json",
    "work_requests": "di_workrequests_latest.json",
    "workrequest": "di_workrequests_latest.json",
    "workrequests": "di_workrequests_latest.json",
    "di": "di_workrequests_latest.json",
}


def cell_to_json(value):
    if value is None:
        return None
    # Keep dates and datetimes readable for SQL import script.
    if hasattr(value, "isoformat"):
        try:
            return value.isoformat()
        except Exception:
            return str(value)
    return value


def iter_sheet_rows(ws) -> Iterable[List[object]]:
    for row in ws.iter_rows(values_only=True):
        yield list(row)


def sheet_to_records(ws) -> List[Dict[str, object]]:
    rows = list(iter_sheet_rows(ws))
    if not rows:
        return []

    headers_raw = rows[0]
    headers = []
    for i, col in enumerate(headers_raw):
        if col is None or str(col).strip() == "":
            headers.append(f"col_{i+1}")
        else:
            headers.append(str(col).strip())

    records: List[Dict[str, object]] = []
    for row in rows[1:]:
        if not any(cell is not None and str(cell).strip() != "" for cell in row):
            continue
        padded = row + [None] * (len(headers) - len(row))
        rec = {headers[i]: cell_to_json(padded[i]) for i in range(len(headers))}
        records.append(rec)
    return records


def output_filename_for_sheet(sheet_name: str) -> str:
    key = normalize_name(sheet_name)
    if key in SHEET_TO_OUTPUT:
        return SHEET_TO_OUTPUT[key]
    return f"{key}.json"


def convert_csv(input_path: Path, out_dir: Path) -> None:
    with input_path.open("r", encoding="utf-8-sig", newline="") as f:
        sample = f.read(4096)
        f.seek(0)
        try:
            dialect = csv.Sniffer().sniff(sample, delimiters=",;\t|")
        except csv.Error:
            dialect = csv.excel
            dialect.delimiter = ";"
        reader = csv.DictReader(f, dialect=dialect)
        records = []
        for row in reader:
            if not any(str(v).strip() for v in row.values() if v is not None):
                continue
            records.append({k: (v if v != "" else None) for k, v in row.items()})

    out_name = output_filename_for_sheet(input_path.stem)
    out_path = out_dir / out_name
    with out_path.open("w", encoding="utf-8") as f:
        json.dump(records, f, ensure_ascii=False, indent=2)

    print("CSV:", input_path)
    print(f"- {input_path.name} -> {out_name} ({len(records)} rows)")


def convert_excel(input_path: Path, out_dir: Path, selected_sheet: Optional[str] = None) -> None:
    try:
        from openpyxl import load_workbook
    except Exception as exc:
        raise RuntimeError(
            "openpyxl is required. Install with: pip install openpyxl"
        ) from exc

    if not input_path.exists():
        raise FileNotFoundError(f"Input file not found: {input_path}")

    out_dir.mkdir(parents=True, exist_ok=True)

    wb = load_workbook(filename=input_path, data_only=True)
    sheets = wb.sheetnames

    if selected_sheet:
        target = None
        wanted = normalize_name(selected_sheet)
        for name in sheets:
            if normalize_name(name) == wanted:
                target = name
                break
        if target is None:
            raise ValueError(f"Sheet not found: {selected_sheet}. Available: {', '.join(sheets)}")
        sheets = [target]

    print("Workbook:", input_path)
    print("Sheets:", ", ".join(sheets))

    for sheet_name in sheets:
        ws = wb[sheet_name]
        records = sheet_to_records(ws)
        out_name = output_filename_for_sheet(sheet_name)
        out_path = out_dir / out_name

        with out_path.open("w", encoding="utf-8") as f:
            json.dump(records, f, ensure_ascii=False, indent=2)

        print(f"- {sheet_name} -> {out_name} ({len(records)} rows)")


def main() -> None:
    parser = argparse.ArgumentParser(description="Convert Excel sheets to JSON files for Docker import.")
    parser.add_argument("--input", required=True, help="Path to .xlsx or .csv file")
    parser.add_argument(
        "--out-dir",
        default=str(Path(__file__).parent / "data_extracted"),
        help="Output directory for json files",
    )
    parser.add_argument("--sheet", default=None, help="Convert one sheet only (optional)")
    args = parser.parse_args()

    input_path = Path(args.input)
    out_dir = Path(args.out_dir)
    out_dir.mkdir(parents=True, exist_ok=True)

    ext = input_path.suffix.lower()
    if ext == ".csv":
        convert_csv(input_path=input_path, out_dir=out_dir)
        return
    if ext in {".xlsx", ".xlsm", ".xltx", ".xltm"}:
        convert_excel(input_path=input_path, out_dir=out_dir, selected_sheet=args.sheet)
        return

    raise ValueError(f"Unsupported input format: {input_path.suffix}. Use .xlsx or .csv")


if __name__ == "__main__":
    main()
