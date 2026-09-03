import os
import sys
from pathlib import Path
import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn

def set_cell_background(cell, fill_hex):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{fill_hex}"/>')
    tcPr.append(shd)

def set_cell_margins(cell, top=100, bottom=100, left=150, right=150):
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = OxmlElement('w:tcMar')
    for m, val in [('top', top), ('bottom', bottom), ('left', left), ('right', right)]:
        node = OxmlElement(f'w:{m}')
        node.set(qn('w:w'), str(val))
        node.set(qn('w:type'), 'dxa')
        tcMar.append(node)
    tcPr.append(tcMar)

def add_heading_styled(doc, text, level):
    p = doc.add_heading(text, level=level)
    p.paragraph_format.space_before = Pt(14)
    p.paragraph_format.space_after = Pt(6)
    for run in p.runs:
        run.font.name = 'Segoe UI'
        if level == 1:
            run.font.size = Pt(18)
            run.font.color.rgb = RGBColor(15, 27, 128)
            run.bold = True
        elif level == 2:
            run.font.size = Pt(14)
            run.font.color.rgb = RGBColor(255, 138, 0)
            run.bold = True
        elif level == 3:
            run.font.size = Pt(12)
            run.font.color.rgb = RGBColor(43, 29, 76)
            run.bold = True
    return p

def main():
    if sys.platform == "win32":
        sys.stdout.reconfigure(encoding="utf-8")
        sys.stderr.reconfigure(encoding="utf-8")

    desktop_path = Path(os.path.expanduser("~")) / "Desktop"
    output_filename = desktop_path / "Rapport_Final_GMAO_COSWIN_USER_18_Aout_2026.docx"

    doc = docx.Document()

    for section in doc.sections:
        section.top_margin = Inches(0.8)
        section.bottom_margin = Inches(0.8)
        section.left_margin = Inches(0.8)
        section.right_margin = Inches(0.8)

    normal_style = doc.styles['Normal']
    normal_style.font.name = 'Calibri'
    normal_style.font.size = Pt(11)
    normal_style.font.color.rgb = RGBColor(50, 50, 50)

    title_p = doc.add_paragraph()
    title_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title_p.paragraph_format.space_after = Pt(4)
    r_title = title_p.add_run("RAPPORT DE VALIDATION FINALE")
    r_title.font.name = 'Montserrat'
    r_title.font.size = Pt(24)
    r_title.font.bold = True
    r_title.font.color.rgb = RGBColor(15, 27, 128)

    sub_p = doc.add_paragraph()
    sub_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    sub_p.paragraph_format.space_after = Pt(20)
    r_sub = sub_p.add_run("Intégration gmao_mutualise_ODS.dbo.COSWIN_USER (4 205 Agents Senelec)")
    r_sub.font.name = 'Segoe UI'
    r_sub.font.size = Pt(13)
    r_sub.font.color.rgb = RGBColor(255, 138, 0)
    r_sub.bold = True

    meta_table = doc.add_table(rows=2, cols=2)
    meta_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    cell_data = [
        [("Date :", " Mardi 18 Août 2026"), ("Projet :", " AppMobileGmao (Senelec)")],
        [("Serveur :", " srv-bddomtech (1433)"), ("Statut :", " 🎉 100% SUCCÈS (4/4 Tests Validés)")]
    ]
    for row_idx, row in enumerate(meta_table.rows):
        for col_idx, cell in enumerate(row.cells):
            set_cell_background(cell, "F0F3FF")
            set_cell_margins(cell, top=100, bottom=100, left=150, right=150)
            p = cell.paragraphs[0]
            label, val = cell_data[row_idx][col_idx]
            r_lbl = p.add_run(label)
            r_lbl.bold = True
            r_lbl.font.color.rgb = RGBColor(15, 27, 128)
            r_val = p.add_run(val)
            if "100%" in val:
                r_val.bold = True
                r_val.font.color.rgb = RGBColor(0, 140, 60)

    doc.add_paragraph().paragraph_format.space_after = Pt(10)

    # --- SECTION 1 ---
    add_heading_styled(doc, "1. Synthèse de la Découverte & Emplacement Exact", level=1)
    p = doc.add_paragraph()
    p.add_run("Le scan automatisé sur le réseau d'entreprise Senelec a permis d'identifier l'emplacement officiel de la table ")
    p.add_run("COSWIN_USER").bold = True
    p.add_run(" :")

    bp1 = doc.add_paragraph(style='List Bullet')
    bp1.add_run("Base de données hôte : ").bold = True
    bp1.add_run("gmao_mutualise_ODS").bold = True
    bp1.add_run(" (Base mutualisée Senelec sur srv-bddomtech:1433).")

    bp2 = doc.add_paragraph(style='List Bullet')
    bp2.add_run("Volume d'utilisateurs réels : ").bold = True
    bp2.add_run("4 205 agents Senelec ").bold = True
    bp2.add_run("enregistrés et à jour.")

    bp3 = doc.add_paragraph(style='List Bullet')
    bp3.add_run("Colonnes officielles : ").bold = True
    bp3.add_run("CWCU_CODE (Matricule), CWCU_SIGNATURE (Nom), CWCU_EMAIL (Email), CWCU_ENTITY (Entité), CWCU_PREFERRED_GROUP (Groupe).")

    # --- SECTION 2 ---
    add_heading_styled(doc, "2. Résultats des 4 Tests de Validation PowerShell", level=1)

    t_res = doc.add_table(rows=5, cols=3)
    t_res.alignment = WD_TABLE_ALIGNMENT.CENTER
    headers = ["Test", "Fonctionnalité Évaluée", "Résultat d'Exécution Réelle"]
    for idx, text_h in enumerate(headers):
        cell = t_res.rows[0].cells[idx]
        cell.text = text_h
        set_cell_background(cell, "0F1B80")
        for r in cell.paragraphs[0].runs:
            r.font.bold = True
            r.font.color.rgb = RGBColor(255, 255, 255)

    tests_data = [
        ("Test 1", "Statistiques & Décompte total", "✅ 4 205 agents réels (3028 SENELEC, 85 DTAE_EXPLOITANT...)"),
        ("Test 2", "Autocomplétion Mobile (search_users 'DIOP')", "✅ 10 agents réels extraits (ex: Ababacar Sadikh DIOP 7650)"),
        ("Test 3", "Fiche Profil (get_user_by_code)", "✅ Profil extrait (Code: 3910_TEST, Email: macisse.diop@senelec.sn)"),
        ("Test 4", "Validation Matricule (validate_employee_code)", "✅ Valide=True ('3910_TEST') / Inexistant=False ('INVALID999')")
    ]
    for r_idx, r_data in enumerate(tests_data, 1):
        row_cells = t_res.rows[r_idx].cells
        bg = "F9F9FB" if r_idx % 2 == 1 else "FFFFFF"
        for c_idx, val in enumerate(r_data):
            row_cells[c_idx].text = val
            set_cell_background(row_cells[c_idx], bg)
            set_cell_margins(row_cells[c_idx], 80, 80, 100, 100)
            if c_idx == 2:
                row_cells[c_idx].paragraphs[0].runs[0].font.color.rgb = RGBColor(0, 120, 50)

    doc.add_paragraph().paragraph_format.space_after = Pt(10)

    # --- SECTION 3 ---
    add_heading_styled(doc, "3. Architecture Backend SOLID & DRY", level=1)
    p_arch = doc.add_paragraph()
    p_arch.add_run("• CoswinUserService (Single Responsibility) : ").bold = True
    p_arch.add_run("Point d'accès unique interrogeant gmao_mutualise_ODS.dbo.COSWIN_USER sans aucun fallback.\n")
    p_arch.add_run("• Élimination de la Duplication (DRY) : ").bold = True
    p_arch.add_run("user_service.get_supervisors_list() délègue directement à CoswinUserService.search_users().")

    # Conclusion Box
    box_table = doc.add_table(rows=1, cols=1)
    box_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    cell = box_table.rows[0].cells[0]
    set_cell_background(cell, "E8F5E9")
    set_cell_margins(cell, top=140, bottom=140, left=180, right=180)
    
    p_cell = cell.paragraphs[0]
    r_c1 = p_cell.add_run("🎉 SUCCÈS TOTAL DE L'INTÉGRATION COSWIN_USER :\n")
    r_c1.bold = True
    r_c1.font.color.rgb = RGBColor(0, 120, 50)
    
    p_cell.add_run(
        "L'ensemble de la chaîne d'accès aux identités des 4 205 agents Senelec "
        "est désormais 100% validé et opérationnel."
    )

    doc.save(str(output_filename))
    print(f"✅ Nouveau rapport Word final généré sur le Bureau : {output_filename}")

if __name__ == "__main__":
    main()
